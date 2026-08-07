"""
Budget vs Actual Data Extraction Script

Extracts financial metrics from Excel files and outputs structured JSON data.
Implements name-based metric searching and automatic month detection.
"""

import pandas as pd
import openpyxl
import json
import argparse
from datetime import datetime
from pathlib import Path


# Metric names to search for (from source file)
METRIC_NAMES = [
    "Total Revenue",
    "Total Variable Cost",
    "Contribution Margin",
    "Total Staff Cost (Direct)",
    "Gross Profit / (Loss)",
    "Total Fixed Cost (Indirect)",
    "Net Surplus / (Deflect)",
    "Interest Expenses",
    "Amortization",
    "Depreciation",
]

# Abbreviated names for percentage metrics (stored in database)
PERCENTAGE_ABBREVIATIONS = {
    "Total Variable Cost": "TVC %",
    "Contribution Margin": "CM %",
    "Total Staff Cost (Direct)": "TSCD %",
    "Gross Profit / (Loss)": "GP %",
    "Total Fixed Cost (Indirect)": "TFCI %",
    "Net Surplus / (Deflect)": "NSD %"
}

# Dashboard display names
DASHBOARD_NAMES = {
    "Total Revenue": "Total Revnue",
    "Total Variable Cost": "Variable Cost",
    "TVC %": "Variable Cost %",
    "Contribution Margin": "Contribution Margin",
    "CM %": "Contribution Margin %",
    "Total Staff Cost (Direct)": "Fixed Costs (Direct)",
    "TSCD %": "Fixed Costs (Direct) %",
    "Gross Profit / (Loss)": "Gross Profit / Loss",
    "GP %": "Gross Profit / Loss %",
    "Total Fixed Cost (Indirect)": "Fixed Costs (Indirect)",
    "TFCI %": "Fixed Costs (Indirect) %",
    "Net Surplus / (Deflect)": "Net Profit / Loss",
    "NSD %": "Net Profit / Loss %",
    "Interest Expenses": "Interest Expenses",
    "Amortization": "Amortization",
    "Depreciation": "Depreciation",
}

def load_company_labels(company):
    config_path = Path(__file__).parent.parent / "config" / "company_labels.json"
    if not config_path.exists():
        return {}
    with open(config_path) as f:
        all_configs = json.load(f)
    return all_configs.get(company, all_configs.get("default", {}))


MONTH_NAMES = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December"
]

_MONTH_NAMES_LOWER = frozenset(m.lower() for m in MONTH_NAMES)


def detect_column_structure(ws):
    """
    Dynamically detect Budget/Actual/Variance column positions for each month.
    Skips YTD/Total summary columns by checking the month label row.
    Returns list of (month_num, budget_col, actual_col, variance_col).
    """
    header_row = None
    for row in range(1, 26):
        for col in range(2, min(ws.max_column + 1, 15)):
            v = ws.cell(row, col).value
            if v and str(v).strip().lower() == "budget":
                header_row = row
                break
        if header_row:
            break

    if not header_row:
        raise ValueError("No 'Budget' header row found in first 25 rows. Check file structure.")

    # Find the row that contains real month names (scan upward from header_row)
    label_row = None
    for row in range(header_row - 1, 0, -1):
        for col in range(1, min(ws.max_column + 1, 20)):
            v = ws.cell(row, col).value
            if v and str(v).strip().lower() in _MONTH_NAMES_LOWER:
                label_row = row
                break
        if label_row:
            break

    months = []
    month_num = 0

    for col in range(2, ws.max_column + 1):
        v = ws.cell(header_row, col).value
        if not v or str(v).strip().lower() != "budget":
            continue

        # Skip Total/YTD columns: check that label row has a real month name at or just before this column
        if label_row is not None:
            lv_exact = ws.cell(label_row, col).value
            lv_prev = ws.cell(label_row, col - 1).value if col > 1 else None
            lv_exact_str = str(lv_exact).strip().lower() if lv_exact else ""
            lv_prev_str = str(lv_prev).strip().lower() if lv_prev else ""
            if lv_exact_str in _MONTH_NAMES_LOWER:
                pass  # Real month at exact col — include
            elif not lv_exact_str and lv_prev_str in _MONTH_NAMES_LOWER:
                pass  # Empty at exact col, real month at col-1 (merged cell) — include
            else:
                continue  # Non-month or YTD/Total label — skip

        budget_col = col
        actual_col = None
        variance_col = None

        for ahead in range(col + 1, min(col + 7, ws.max_column + 2)):
            av = ws.cell(header_row, ahead).value
            if av is None:
                continue
            av_str = str(av).strip().lower()
            # Tolerant prefix matching for typos like "Acutal" / "Varaince"
            if av_str[:2] == "ac" and actual_col is None:
                actual_col = ahead
            elif av_str[:3] == "var" and actual_col is not None:
                variance_col = ahead
                break
            elif av_str == "budget":
                break

        if actual_col and variance_col:
            month_num += 1
            months.append((month_num, budget_col, actual_col, variance_col))

    if not months:
        raise ValueError("No complete Budget/Actual/Variance column groups found.")

    return months


def detect_month_from_columns(ws, budget_header_row=11):
    """Detect number of months in the file using dynamic column detection."""
    col_structure = detect_column_structure(ws)
    month_count = len(col_structure)
    if month_count < 1:
        raise ValueError("No 'Budget' headers found in file. Check file structure.")
    month_name = MONTH_NAMES[month_count - 1]
    print(f"Found {month_count} month(s) of data")
    return month_count, month_name


def find_metric_row(ws, metric_name, search_col=1, max_row=150):
    """Find the row number where a metric name appears. Exact match takes priority over startswith."""
    first_startswith = None
    for row in range(1, max_row + 1):
        cell_value = ws.cell(row, search_col).value
        if cell_value:
            cell_str = str(cell_value).strip()
            if cell_str == metric_name:
                return row
            if first_startswith is None and cell_str.startswith(metric_name):
                first_startswith = row
    return first_startswith


def extract_metric_data(ws, metric_name, row_num, month_number, col_structure=None):
    """
    Extract Budget, Actual, Variance for a specific metric and month.

    Each month has 3 columns: Budget, Actual, Variance
    Column A = labels (offset 1)
    Month 1 (Jan) = columns 2, 3, 4 (Budget, Actual, Variance)
    Month 2 (Feb) = columns 5, 6, 7
    Month N = columns (N*3 - 1), (N*3), (N*3 + 1)

    Args:
        ws: openpyxl worksheet
        metric_name: Name of the metric
        row_num: Row number where metric is located
        month_number: Which month to extract (1-12)
        col_structure: Optional list from detect_column_structure()

    Returns:
        dict: {budget, actual, variance}
    """
    # Calculate column positions for this month
    if col_structure:
        month_entry = next((m for m in col_structure if m[0] == month_number), None)
        if not month_entry:
            return {"budget": 0.0, "actual": 0.0, "variance": 0.0}
        _, budget_col, actual_col, variance_col = month_entry
    else:
        budget_col = (month_number * 3) - 1
        actual_col = month_number * 3
        variance_col = (month_number * 3) + 1

    budget = ws.cell(row_num, budget_col).value
    actual = ws.cell(row_num, actual_col).value
    variance = ws.cell(row_num, variance_col).value

    # Convert to float, handle None values and Excel errors
    def safe_float(value):
        if value is None:
            return 0.0
        if isinstance(value, str) and ('#' in value or value.strip() == ''):
            return 0.0
        try:
            return float(value)
        except (ValueError, TypeError):
            return 0.0

    budget = safe_float(budget)
    actual = safe_float(actual)
    variance = safe_float(variance)

    return {
        "budget": budget,
        "actual": actual,
        "variance": variance
    }


def extract_all_metrics(file_path, sheet_name="Detail Budget", extract_all_months=True, company="default"):
    """
    Extract all metrics from the Excel file.

    Args:
        file_path: Path to Excel file
        sheet_name: Name of sheet to extract from
        extract_all_months: If True, extract all months; if False, only extract last month
        company: Company ID for label overrides

    Returns:
        list: List of metric dictionaries in RAW DATA format
    """
    print(f"Loading workbook: {file_path}")
    wb = openpyxl.load_workbook(file_path, data_only=True)

    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found. Available sheets: {wb.sheetnames}")

    ws = wb[sheet_name]

    # Detect total number of months in file
    col_structure = detect_column_structure(ws)
    total_months = len(col_structure)
    last_month_name = MONTH_NAMES[total_months - 1]
    print(f"Detected {total_months} month(s) of data (last month: {last_month_name})")

    # Build company-specific label lookup: standard_key -> search_label
    company_config = load_company_labels(company)
    label_overrides = company_config.get("bva_row_labels", {})
    label_col = company_config.get("label_col", 1)
    # search_labels[standard_key] = label to search for in file (None = derived)
    search_labels = {m: label_overrides.get(m, m) for m in METRIC_NAMES}

    # Determine which months to extract
    if extract_all_months:
        months_to_extract = range(1, total_months + 1)
        print(f"Extracting ALL {total_months} months")
    else:
        months_to_extract = [total_months]
        print(f"Extracting only last month: {last_month_name}")

    extracted_data = []
    monthly_extracted = {}

    # Extract each month
    for month_number in months_to_extract:
        month_name = MONTH_NAMES[month_number - 1]
        print(f"\n--- Processing {month_name} (Month {month_number}) ---")

        # Extract each metric for this month
        for metric_name in METRIC_NAMES:
            search_label = search_labels[metric_name]

            # Skip derived metrics (search_label is None)
            if search_label is None:
                print(f"  Skipping derived metric: {metric_name}")
                continue

            print(f"  Searching for metric: {metric_name} (label: '{search_label}')")

            # Find the metric row using the company-specific label
            row_num = find_metric_row(ws, search_label, search_col=label_col)

            if row_num is None:
                print(f"    WARNING: Metric '{search_label}' not found. Skipping.")
                continue

            print(f"    [OK] Found at row {row_num}")

            # Extract the metric data for this specific month
            data = extract_metric_data(ws, search_label, row_num, month_number, col_structure)

            # Track for derived metric computation (always keyed by standard name)
            if month_name not in monthly_extracted:
                monthly_extracted[month_name] = {}
            monthly_extracted[month_name][metric_name] = data

            # Add to results — always store standard key as DataPoint
            extracted_data.append({
                "Month": month_name,
                "DataPoint": metric_name,
                "DashboardName": DASHBOARD_NAMES.get(metric_name, metric_name),
                "Budget": data["budget"],
                "Actual": data["actual"],
                "Variance": data["variance"]
            })

            # Check for percentage in row below
            if metric_name in PERCENTAGE_ABBREVIATIONS:
                percentage_row = row_num + 1
                percentage_cell = ws.cell(percentage_row, 1).value

                # Check if next row contains "%"
                if percentage_cell and "%" in str(percentage_cell):
                    print(f"    [OK] Found percentage at row {percentage_row}")

                    # Extract percentage data
                    pct_data = extract_metric_data(ws, f"{search_label} %", percentage_row, month_number, col_structure)
                    pct_abbrev = PERCENTAGE_ABBREVIATIONS[metric_name]

                    extracted_data.append({
                        "Month": month_name,
                        "DataPoint": pct_abbrev,
                        "DashboardName": DASHBOARD_NAMES.get(pct_abbrev, pct_abbrev),
                        "Budget": pct_data["budget"],
                        "Actual": pct_data["actual"],
                        "Variance": pct_data["variance"]
                    })

    # Second pass: compute derived metrics (e.g. Gross Profit when not in source file)
    for month_number in (months_to_extract if not isinstance(months_to_extract, range) else list(months_to_extract)):
        month_name = MONTH_NAMES[month_number - 1]
        mv = monthly_extracted.get(month_name, {})
        for metric_name in METRIC_NAMES:
            if search_labels[metric_name] is not None:
                continue
            if metric_name == "Gross Profit / (Loss)":
                cm = mv.get("Contribution Margin")
                sc = mv.get("Total Staff Cost (Direct)")
                if not cm or not sc:
                    print(f"    WARNING: Cannot derive '{metric_name}' — missing CM or Staff Cost data")
                    continue
                gp_b = round(cm['budget'] - sc['budget'], 2)
                gp_a = round(cm['actual'] - sc['actual'], 2)
                gp_v = round(gp_a - gp_b, 2)
                extracted_data.append({
                    "Month": month_name,
                    "DataPoint": metric_name,
                    "DashboardName": DASHBOARD_NAMES.get(metric_name, metric_name),
                    "Budget": gp_b,
                    "Actual": gp_a,
                    "Variance": gp_v
                })
                print(f"    [DERIVED] '{metric_name}' = CM - Staff Cost")
                # Derive GP %
                if metric_name in PERCENTAGE_ABBREVIATIONS:
                    rev = mv.get("Total Revenue")
                    rev_b = rev['budget'] if rev and rev['budget'] else 0
                    rev_a = rev['actual'] if rev and rev['actual'] else 0
                    gp_pct_b = round(gp_b / rev_b, 6) if rev_b else 0.0
                    gp_pct_a = round(gp_a / rev_a, 6) if rev_a else 0.0
                    pct_abbrev = PERCENTAGE_ABBREVIATIONS[metric_name]
                    extracted_data.append({
                        "Month": month_name,
                        "DataPoint": pct_abbrev,
                        "DashboardName": DASHBOARD_NAMES.get(pct_abbrev, pct_abbrev),
                        "Budget": gp_pct_b,
                        "Actual": gp_pct_a,
                        "Variance": round(gp_pct_a - gp_pct_b, 6)
                    })

    wb.close()

    unique_months = sorted(set(record['Month'] for record in extracted_data))
    print(f"\n[OK] Extracted {len(extracted_data)} total records for {len(unique_months)} month(s): {', '.join(unique_months)}")

    return extracted_data


def save_to_json(data, output_path):
    """
    Save extracted data to JSON file.

    Args:
        data: List of metric dictionaries
        output_path: Path to output JSON file
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with open(output_path, 'w') as f:
        json.dump(data, f, indent=2)

    print(f"\n[OK] Data saved to: {output_path}")


def main():
    parser = argparse.ArgumentParser(
        description="Extract Budget vs Actual data from Excel files"
    )
    parser.add_argument(
        "--input",
        required=True,
        help="Path to input Excel file"
    )
    parser.add_argument(
        '--company',
        required=True,
        help='Company ID (e.g. mudin, mantis)'
    )
    parser.add_argument(
        "--output",
        default=None,
        help="Path to output JSON file"
    )
    parser.add_argument(
        "--sheet",
        default="Detail Budget",
        help="Sheet name to extract from (default: Detail Budget)"
    )
    parser.add_argument(
        "--all-months",
        action="store_true",
        default=True,
        help="Extract all months (default: True)"
    )
    parser.add_argument(
        "--last-month-only",
        action="store_true",
        help="Extract only the last month (overrides --all-months)"
    )

    args = parser.parse_args()

    if args.output is None:
        args.output = f"dashboard/data/raw-data-{args.company}.json"

    print("="*60)
    print("BUDGET VS ACTUAL DATA EXTRACTION")
    print("="*60)

    try:
        # Use company-specific source sheet if configured
        company_config = load_company_labels(args.company)
        sheet_to_use = company_config.get("source_sheet", args.sheet)
        if sheet_to_use != args.sheet:
            print(f"Using company-configured source sheet: {sheet_to_use}")

        # Determine extraction mode
        extract_all = not args.last_month_only

        # Extract data
        data = extract_all_metrics(args.input, sheet_to_use, extract_all_months=extract_all, company=args.company)

        # Save to JSON
        save_to_json(data, args.output)

        print("\n" + "="*60)
        print("[SUCCESS] EXTRACTION COMPLETE")
        print("="*60)

    except Exception as e:
        print(f"\n[ERROR] {e}")
        raise


if __name__ == "__main__":
    main()
