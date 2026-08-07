"""
Shared xlsx utilities. safe_load_workbook() wraps openpyxl.load_workbook
and automatically patches known stylesheet issues (e.g. font family
values > 14, which some third-party Excel generators produce) before
retrying, so the rest of the pipeline never sees the error.
"""
import re
import tempfile
import zipfile
from pathlib import Path

import openpyxl


def _sanitize_to_temp(file_path):
    """
    Write a patched copy of the xlsx to a system temp file and return its path.
    Never touches the original — avoids file-locking issues (e.g. OneDrive sync).
    Caller is responsible for deleting the temp file when done.
    """
    file_path = Path(file_path)
    tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
    tmp.close()
    tmp_path = Path(tmp.name)
    try:
        with zipfile.ZipFile(file_path, 'r') as zin:
            with zipfile.ZipFile(tmp_path, 'w', zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if item.filename == 'xl/styles.xml':
                        xml_str = data.decode('utf-8')
                        xml_str = re.sub(
                            r'<family val="(\d+)"/>',
                            lambda m: f'<family val="{min(int(m.group(1)), 14)}"/>',
                            xml_str
                        )
                        data = xml_str.encode('utf-8')
                    zout.writestr(item, data)
        return tmp_path
    except Exception as e:
        tmp_path.unlink(missing_ok=True)
        raise RuntimeError(f"Could not patch stylesheet in {file_path.name}: {e}") from e


def safe_load_workbook(file_path, **kwargs):
    """
    Load an xlsx workbook, patching stylesheet errors once and retrying.
    Drop-in replacement for openpyxl.load_workbook.
    """
    try:
        return openpyxl.load_workbook(file_path, **kwargs)
    except ValueError as exc:
        if 'stylesheet' not in str(exc).lower():
            raise
        print(f"  [INFO] Patching invalid stylesheet in {Path(file_path).name} ...")
        tmp_path = _sanitize_to_temp(file_path)
        try:
            return openpyxl.load_workbook(tmp_path, **kwargs)
        finally:
            tmp_path.unlink(missing_ok=True)
