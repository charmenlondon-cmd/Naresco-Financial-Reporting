"""
Section Detail Data Extractor

Extracts named row(s) from the BVA source file (Budget/Actual/Variance per month)
and writes them into a dedicated sheet in the company database.
Designed to be extended: add new sections to SECTIONS_CONFIG as needed.

DB sheet structure per section:
  Description | Jan-26 Budget | Jan-26 Actual | Jan-26 Variance | Feb-26 Budget | ... | YTD Budget | YTD Actual | YTD Variance
"""

import openpyxl
from openpyxl import load_workbook, Workbook
import argparse
from pathlib import Path
from datetime import datetime


MONTH_NAMES = [
    "January","February","March","April","May","June",
    "July","August","September","October","November","December"
]

# Config: each section maps to the row labels to extract from the BVA source
SECTIONS_CONFIG = {
    "Variable Costs": {
        "metrics": ["Material Cost - Projects"],
        "total_label": "TOTAL VARIABLE COSTS",
        "source_sheet": "Detail Budget"
    },
    "Fixed Costs": {
        "metrics": [
            "Staff Salaries",
            "Labor Salaries",
            "Wages & Salaries",
            "Recruitment Expenses",
            "Staff Gratuity",
            "Staff Air Ticket Provision",
            "Staff Visa Cost",
            "Staff Insurance Expense",
            "Staff Accommodation Expense",
            "Staff Welfare",
            "General Insurance",
            "Pension Expense 12.5%",
        ],
        "total_label": "TOTAL FIXED COSTS",
        "source_sheet": "Detail Budget"
    },
    "Indirect Costs": {
        "metrics": [
            "Director Remuneration",
            "Office Rent",
            "YARD Expenses",
            "DEWA (Electricity & Water) - Office",
            "Facility Rent",
            "Telephone & Internet",
            "Travelling & Transport - Admin",
            "Printing & Stationery",
            "Building Insurance",
            "Building Maintenance",
            "Repair & Maintenance - Vehicle",
            "Repair & Maintenance - Machines",
            "Group Fee",
            "AAAI Management Fee",
            "Divisional Salaries & Other Expenses",
            "Administration Charges",
            "HR Service charges",
            "IT Expenses",
            "Software Charges",
            "Fine & Penalty",
            "VAT Expenses",
            "Audit fee",
            "Tender Fees",
            "Trade License Fee",
            "Government & Municipality Charges",
            "Bank Charges",
            "Donation",
            "Business Promotion Expense",
            "Advertising & Social Media Expenses",
            "Travelling & Transport - Marketing",
        ],
        "total_label": "TOTAL INDIRECT COSTS",
        "source_sheet": "Detail Budget"
    },
    "TDI": {
        "metrics": [
            "Interest Expenses",
            "Amortization",
            "Depreciation",
        ],
        "total_label": "TOTAL TDI",
        "source_sheet": "Detail Budget"
    },
}


def detect_column_structure(ws):
    """
    Dynamically detect Budget/Actual/Variance column positions for each month.
    Returns list of (month_num, month_label, budget_col, actual_col, variance_col).
    """
    header_row = None
    for row in range(1, 26):
        for col in range(2, min(ws.max_column + 1, 15)):
            v = ws.cell(row, col).value
            if v and str(v).strip().lower() == "budget":
                header_row = row
                break
        if header_row:
            break

    if not header_row:
        raise ValueError("No 'Budget' header row found in first 25 rows.")

    months = []
    month_num = 0

    for col in range(2, ws.max_column + 1):
        v = ws.cell(header_row, col).value
        if not v or str(v).strip().lower() != "budget":
            continue

        budget_col = col
        actual_col = None
        variance_col = None

        for ahead in range(col + 1, min(col + 7, ws.max_column + 2)):
            av = ws.cell(header_row, ahead).value
            if av is None:
                continue
            av_str = str(av).strip().lower()
            if "actual" in av_str and actual_col is None:
                actual_col = ahead
            elif "variance" in av_str and actual_col is not None:
                variance_col = ahead
                break
            elif av_str == "budget":
                break

        if actual_col and variance_col:
            month_num += 1
            label = MONTH_NAMES[month_num - 1][:3] + "-26"
            months.append((month_num, label, budget_col, actual_col, variance_col))

    return months


def find_metric_row(ws, metric_name, search_col=1, max_row=200):
    """Find row by label (exact or startswith match)."""
    for row in range(1, max_row + 1):
        v = ws.cell(row, search_col).value
        if v:
            s = str(v).strip()
            if s == metric_name or s.startswith(metric_name):
                return row
    return None


def safe_float(v):
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    return 0.0


def extract_section_from_source(ws, metrics, col_structure):
    """
    Extract Budget/Actual/Variance for listed metric rows across all months.
    Returns list of dicts: {description, month_label: {budget, actual, variance}}
    """
    results = []
    for metric in metrics:
        row_num = find_metric_row(ws, metric)
        if row_num is None:
            print(f"    WARNING: '{metric}' not found — skipping")
            continue

        row_data = {"description": metric}
        for _, label, bcol, acol, vcol in col_structure:
            row_data[label] = {
                "budget":   safe_float(ws.cell(row_num, bcol).value),
                "actual":   safe_float(ws.cell(row_num, acol).value),
                "variance": safe_float(ws.cell(row_num, vcol).value),
            }
        results.append(row_data)
        print(f"    [OK] '{metric}' extracted")

    return results


def read_existing_section_sheet(ws):
    """
    Read existing section sheet from database.
    Returns: (month_labels, existing_rows)
    month_labels: list of month label strings already in DB
    existing_rows: list of dicts matching extract_section_from_source output
    """
    if ws.max_row < 1:
        return [], []

    # Row 1 is header: Description | Jan-26 Budget | Jan-26 Actual | Jan-26 Variance | ... | YTD Budget | ...
    headers = []
    for col in range(1, ws.max_column + 1):
        v = ws.cell(1, col).value
        if v is None:
            break
        headers.append(v)

    # Find month labels (headers that end in " Budget" and aren't "YTD Budget")
    month_labels = []
    seen = set()
    for h in headers:
        if h and h.endswith(" Budget") and not h.startswith("YTD"):
            label = h[:-7]  # strip " Budget"
            if label not in seen:
                seen.add(label)
                month_labels.append(label)

    existing_rows = []
    for row in range(2, ws.max_row + 1):
        desc = ws.cell(row, 1).value
        if not desc:
            continue

        row_data = {"description": str(desc).strip()}
        col_offset = 2
        for label in month_labels:
            row_data[label] = {
                "budget":   safe_float(ws.cell(row, col_offset).value),
                "actual":   safe_float(ws.cell(row, col_offset + 1).value),
                "variance": safe_float(ws.cell(row, col_offset + 2).value),
            }
            col_offset += 3
        existing_rows.append(row_data)

    return month_labels, existing_rows


def write_section_sheet(ws, all_months, all_rows, total_label):
    """
    Write/rewrite the section sheet with all months + YTD columns.
    all_months: list of month label strings in chronological order
    all_rows: list of row dicts (excluding total row)
    """
    ws.delete_rows(1, ws.max_row)

    # Header row
    headers = ["Description"]
    for label in all_months:
        headers += [f"{label} Budget", f"{label} Actual", f"{label} Variance"]
    headers += ["YTD Budget", "YTD Actual", "YTD Variance"]

    for col, h in enumerate(headers, 1):
        ws.cell(1, col, h)

    # Data rows
    def write_row(ws_row, row_data):
        ws.cell(ws_row, 1, row_data["description"])
        col = 2
        ytd_b = ytd_a = ytd_v = 0.0
        for label in all_months:
            vals = row_data.get(label, {"budget": 0.0, "actual": 0.0, "variance": 0.0})
            ws.cell(ws_row, col,     vals["budget"])
            ws.cell(ws_row, col + 1, vals["actual"])
            ws.cell(ws_row, col + 2, vals["variance"])
            ytd_b += vals["budget"]
            ytd_a += vals["actual"]
            ytd_v += vals["variance"]
            col += 3
        ws.cell(ws_row, col,     round(ytd_b, 2))
        ws.cell(ws_row, col + 1, round(ytd_a, 2))
        ws.cell(ws_row, col + 2, round(ytd_v, 2))

    for r_idx, row_data in enumerate(all_rows, start=2):
        write_row(r_idx, row_data)

    # Totals row
    total_row = 2 + len(all_rows)
    totals = {"description": total_label}
    for label in all_months:
        totals[label] = {
            "budget":   round(sum(r.get(label, {}).get("budget",   0.0) for r in all_rows), 2),
            "actual":   round(sum(r.get(label, {}).get("actual",   0.0) for r in all_rows), 2),
            "variance": round(sum(r.get(label, {}).get("variance", 0.0) for r in all_rows), 2),
        }
    write_row(total_row, totals)

    print(f"    Written {len(all_rows)} row(s) + totals row, months: {all_months}")


def update_section_in_database(database_path, section_name, section_config, source_rows, source_months):
    """Open/create database and update the named section sheet."""
    database_path = Path(database_path)

    if database_path.exists():
        wb = load_workbook(database_path)
    else:
        wb = Workbook()
        wb.active.title = "Sheet1"

    sheet_name = section_name  # e.g. "Variable Costs"

    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        existing_months, existing_rows = read_existing_section_sheet(ws)
    else:
        ws = wb.create_sheet(sheet_name)
        existing_months = []
        existing_rows = []

    new_months = [m for m in source_months if m not in existing_months]
    all_months = existing_months + new_months

    if not new_months:
        print(f"    No new months — rebuilding rows from current config")
    else:
        print(f"    Adding month(s): {new_months}")

    # Build lookup dicts
    existing_by_desc = {r["description"]: r for r in existing_rows}
    source_by_desc   = {r["description"]: r for r in source_rows}

    # Always build from CURRENT CONFIG metrics only — self-healing when config changes
    all_rows = []
    for metric in section_config["metrics"]:
        row_data = {"description": metric}
        for label in existing_months:
            row_data[label] = existing_by_desc.get(metric, {}).get(label, {"budget":0.0,"actual":0.0,"variance":0.0})
        for label in new_months:
            row_data[label] = source_by_desc.get(metric, {}).get(label, {"budget":0.0,"actual":0.0,"variance":0.0})
        all_rows.append(row_data)
    write_section_sheet(ws, all_months, all_rows, section_config["total_label"])

    wb.save(database_path)
    wb.close()


def main():
    parser = argparse.ArgumentParser(description="Extract BVA section detail into company database")
    parser.add_argument("--input", required=True, help="Path to source BVA Excel file")
    parser.add_argument("--company", required=True, help="Company ID (e.g. mantis)")
    parser.add_argument("--database", default=None)
    args = parser.parse_args()

    if args.database is None:
        args.database = f"Financial-Data-Database-{args.company.capitalize()}.xlsx"

    print("=" * 60)
    print("SECTION DETAIL EXTRACTION")
    print("=" * 60)
    print(f"Source:   {args.input}")
    print(f"Database: {args.database}")

    try:
        wb = openpyxl.load_workbook(args.input, data_only=True)

        for section_name, config in SECTIONS_CONFIG.items():
            print(f"\n--- Section: {section_name} ---")
            sheet_name = config["source_sheet"]
            if sheet_name not in wb.sheetnames:
                print(f"  WARNING: Sheet '{sheet_name}' not found — skipping")
                continue

            ws = wb[sheet_name]
            col_structure = detect_column_structure(ws)
            source_months = [label for _, label, _, _, _ in col_structure]
            print(f"  Source months detected: {source_months}")

            source_rows = extract_section_from_source(ws, config["metrics"], col_structure)
            if not source_rows:
                print(f"  No rows extracted for '{section_name}' — skipping")
                continue

            update_section_in_database(args.database, section_name, config, source_rows, source_months)

        wb.close()

        print()
        print("=" * 60)
        print("[SUCCESS] SECTION EXTRACTION COMPLETE")
        print("=" * 60)

    except Exception as e:
        print(f"\n[ERROR] {e}")
        raise


if __name__ == "__main__":
    main()
